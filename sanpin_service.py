# -*- coding: utf-8 -*-
"""
Сервис проверки меню (ежедневного и типового) на соответствие нормам СанПиН
с приоритетом практических правил, выведенных из эталонных файлов /food.
"""
from __future__ import annotations

import json
import re
import statistics
import time
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
CACHE_PATH = BASE_DIR / "data" / "sanpin_practical_rules.json"
TEMP_CHECK_DIR = BASE_DIR / "temp_sanpin_checks"

# Сколько свежих ежедневных меню брать с каждой школы при построении эталона
DAILY_SAMPLE_PER_SCHOOL = 2
# Минимальное число наблюдений, чтобы считать практическое правило надёжным
MIN_PRACTICAL_SAMPLES = 5
# Калорийность: мелкий разброс рецептур и округления — не ошибка
CALORIE_ABS_EPSILON = 5.0  # ккал
# Ошибка только при заметном занижении относительно эталонного минимума
CALORIE_REL_FLOOR = 0.85


# ---------------------------------------------------------------------------
# Модели
# ---------------------------------------------------------------------------

@dataclass
class MenuRow:
    row_number: int  # 1-based Excel row
    meal: str  # Приём пищи (ПП)
    section: str  # Раздел
    name: str  # Наименование / Блюдо
    exit_g: Optional[float]
    protein: Optional[float]
    fat: Optional[float]
    carb: Optional[float]
    calories: Optional[float]


@dataclass
class PracticalSectionRule:
    section: str
    exit_min: float
    exit_max: float
    exit_avg: float
    cal_avg: Optional[float]
    cal_min: Optional[float]
    prot_avg: Optional[float]
    fat_avg: Optional[float]
    carb_avg: Optional[float]
    meals: list[str]
    samples: int


@dataclass
class DishNorm:
    name: str
    cal_avg: float
    cal_min: float
    exit_avg: float
    prot_avg: Optional[float]
    fat_avg: Optional[float]
    carb_avg: Optional[float]
    samples: int


@dataclass
class PracticalRules:
    sections: dict[str, PracticalSectionRule]
    dishes: dict[str, DishNorm]
    built_at: str
    files_scanned: int
    schools_scanned: int


@dataclass
class Violation:
    rule_id: str
    rule_source: str  # practical | sanpin
    rule_title: str
    row_number: int
    dish_name: str
    section: str
    meal: str
    field: str
    current_value: Any
    recommended_value: Any
    suggestion: str
    kbju_suggestion: Optional[dict[str, float]] = None


@dataclass
class FileCheckResult:
    filename: str
    file_type: str  # daily | typovoye | unknown
    rows_checked: int
    violations: list[Violation] = field(default_factory=list)
    error: Optional[str] = None


# ---------------------------------------------------------------------------
# СанПиН — запасные правила (каждое условие отдельно)
# ---------------------------------------------------------------------------

def _norm_section(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace("ё", "е")


def _norm_meal(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace("ё", "е")


def _norm_dish(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("ё", "е")
    text = re.sub(r"[№#]\s*\d+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _is_calorie_too_low(current: float, norm: float) -> bool:
    """True только при существенно заниженной калорийности (не при 276 vs 287)."""
    if norm <= 0:
        return False
    if current + CALORIE_ABS_EPSILON >= norm:
        return False
    if current >= norm * CALORIE_REL_FLOOR:
        return False
    return True


def _to_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    text = re.sub(r"[^\d.\-]", "", text)
    if not text or text in {".", "-", "-."}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _avg(values: list[float]) -> Optional[float]:
    return round(statistics.mean(values), 2) if values else None


SANPIN_FALLBACK_RULES: list[dict[str, Any]] = [
    {
        "id": "S1",
        "title": "Закуска: выход меньше 60 г",
        "field": "Выход (г)",
        "section_eq": "закуска",
        "exit_lt": 60,
        "exit_gt": 0,
        "recommend": 60,
    },
    {
        "id": "S2",
        "title": "1 блюдо: выход меньше 200 г",
        "field": "Выход (г)",
        "section_eq": "1 блюдо",
        "exit_lt": 200,
        "exit_gt": 0,
        "recommend": 200,
    },
    {
        "id": "S3",
        "title": "2 блюдо: выход меньше 90 г",
        "field": "Выход (г)",
        "section_eq": "2 блюдо",
        "exit_lt": 90,
        "exit_gt": 0,
        "recommend": 90,
    },
    {
        "id": "S4",
        "title": "Гарнир: выход меньше 150 г",
        "field": "Выход (г)",
        "section_eq": "гарнир",
        "exit_lt": 150,
        "exit_gt": 0,
        "recommend": 150,
    },
    {
        "id": "S5",
        "title": "Гор.блюдо: выход меньше 150 г",
        "field": "Выход (г)",
        "section_eq": "гор.блюдо",
        "exit_lt": 150,
        "exit_gt": 0,
        "recommend": 150,
    },
    # S6 и S7 — калорийность; обрабатываются отдельно (справочник)
    {
        "id": "S8",
        "title": "Напиток/сладкое на обед: выход меньше 180 г",
        "field": "Выход (г)",
        "exit_lt": 180,
        "special": "drink_or_sweet_lunch",
        "recommend": 180,
    },
    # S9 — сумма БЖУ; отдельно
    {
        "id": "S10",
        "title": "Выход меньше 200 г (сценарий 10)",
        "field": "Выход (г)",
        "exit_lt": 200,
        "exit_gt": 0,
        "recommend": 200,
        "global_min": True,
    },
    {
        "id": "S11",
        "title": "Выход меньше 200 г (сценарий 11)",
        "field": "Выход (г)",
        "exit_lt": 200,
        "exit_gt": 0,
        "recommend": 200,
        "global_min": True,
    },
    {
        "id": "S12",
        "title": "Выход меньше 500 г (сценарий 12)",
        "field": "Выход (г)",
        "exit_lt": 500,
        "exit_gt": 0,
        "recommend": 500,
        "global_min": True,
    },
    {
        "id": "S13",
        "title": "Выход меньше 700 г (сценарий 13)",
        "field": "Выход (г)",
        "exit_lt": 700,
        "exit_gt": 0,
        "recommend": 700,
        "global_min": True,
    },
    {
        "id": "S14",
        "title": "Выход меньше 300 г (сценарий 14)",
        "field": "Выход (г)",
        "exit_lt": 300,
        "exit_gt": 0,
        "recommend": 300,
        "global_min": True,
    },
    {
        "id": "S15",
        "title": "Выход меньше 500 г (сценарий 15)",
        "field": "Выход (г)",
        "exit_lt": 500,
        "exit_gt": 0,
        "recommend": 500,
        "global_min": True,
    },
    {
        "id": "S16",
        "title": "Фрукты: выход меньше 100 г",
        "field": "Выход (г)",
        "section_startswith": "фрукт",
        "exit_lt": 100,
        "exit_gt": 0,
        "recommend": 100,
    },
]


# ---------------------------------------------------------------------------
# Разбор Excel
# ---------------------------------------------------------------------------

_SKIP_SECTION_MARKERS = ("итого", "среднее", "всего", "сумма")


def _find_header_row(rows: list[tuple]) -> Optional[int]:
    keys = ("раздел", "блюдо", "наименование", "выход", "вес блюда", "калорийность", "прием пищи")
    for i, row in enumerate(rows[:15]):
        if not row:
            continue
        joined = " | ".join(str(c).lower() for c in row if c is not None)
        hits = sum(1 for k in keys if k in joined)
        if hits >= 2 and ("раздел" in joined or "блюд" in joined):
            return i
    return None


def _map_columns(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(headers):
        h = (raw or "").strip().lower().replace("ё", "е")
        if not h:
            continue
        if "прием пищи" in h or h == "пп":
            mapping.setdefault("meal", idx)
        elif "раздел" in h:
            mapping.setdefault("section", idx)
        elif h in {"блюдо", "наименование"} or "блюд" in h and "вес" not in h:
            mapping.setdefault("name", idx)
        elif "выход" in h or "вес блюда" in h or h.startswith("вес"):
            mapping.setdefault("exit", idx)
        elif "белк" in h:
            mapping.setdefault("protein", idx)
        elif "жир" in h:
            mapping.setdefault("fat", idx)
        elif "углевод" in h:
            mapping.setdefault("carb", idx)
        elif "калор" in h:
            mapping.setdefault("calories", idx)
    return mapping


def detect_file_type(filename: str) -> str:
    name = filename.lower()
    if re.match(r"^tm\d{4}-sm\.xlsx?$", name) or "типов" in name:
        return "typovoye"
    if re.match(r"^\d{4}-\d{2}-\d{2}-sm\.xlsx?$", name):
        return "daily"
    return "unknown"


def parse_menu_file(path: Path) -> tuple[list[MenuRow], str]:
    """Читает меню из .xlsx. Возвращает (строки, тип файла)."""
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ValueError(
            "Формат .xls не поддерживается напрямую. Сохраните файл как .xlsx и загрузите снова."
        )
    if suffix not in {".xlsx"}:
        raise ValueError(f"Неподдерживаемый формат: {suffix}. Ожидается .xlsx")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return [], detect_file_type(path.name)

    header_idx = _find_header_row(rows)
    if header_idx is None:
        raise ValueError("Не найдена строка заголовков (Раздел / Блюдо / Выход)")

    headers = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
    cols = _map_columns(headers)
    if "section" not in cols and "name" not in cols:
        raise ValueError("В файле нет колонок «Раздел» или «Блюдо/Наименование»")

    file_type = detect_file_type(path.name)
    current_meal = ""
    result: list[MenuRow] = []

    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(key: str) -> Any:
            idx = cols.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        meal_raw = cell("meal")
        if meal_raw not in (None, ""):
            current_meal = str(meal_raw).strip()

        section_raw = cell("section")
        section = str(section_raw).strip() if section_raw not in (None, "") else ""
        section_l = _norm_section(section)
        if any(m in section_l for m in _SKIP_SECTION_MARKERS):
            continue

        name_raw = cell("name")
        name = str(name_raw).strip() if name_raw not in (None, "") else ""
        if not name and not section:
            continue
        # Строки-итоги по наименованию
        name_l = name.lower()
        if any(m in name_l for m in _SKIP_SECTION_MARKERS):
            continue

        result.append(
            MenuRow(
                row_number=offset,
                meal=current_meal,
                section=section,
                name=name or "—",
                exit_g=_to_float(cell("exit")),
                protein=_to_float(cell("protein")),
                fat=_to_float(cell("fat")),
                carb=_to_float(cell("carb")),
                calories=_to_float(cell("calories")),
            )
        )

    return result, file_type


# ---------------------------------------------------------------------------
# Построение практических правил из /food
# ---------------------------------------------------------------------------

def _iter_reference_files(base_dir: Path = BASE_DIR) -> Iterable[Path]:
    schools = sorted(
        (p for p in base_dir.iterdir() if p.is_dir() and p.name.isdigit()),
        key=lambda p: int(p.name),
    )
    for school in schools:
        food = school / "food"
        if not food.is_dir():
            continue
        # Типовое меню — эталон высокого приоритета
        for tm in sorted(food.glob("tm*-sm.xlsx")):
            yield tm
        dailies = sorted(
            f for f in food.glob("*-sm.xlsx") if not f.name.lower().startswith("tm")
        )
        for daily in dailies[-DAILY_SAMPLE_PER_SCHOOL:]:
            yield daily


def build_practical_rules(
    base_dir: Path = BASE_DIR,
    max_files: Optional[int] = None,
) -> PracticalRules:
    section_exit: dict[str, list[float]] = {}
    section_cal: dict[str, list[float]] = {}
    section_prot: dict[str, list[float]] = {}
    section_fat: dict[str, list[float]] = {}
    section_carb: dict[str, list[float]] = {}
    section_meals: dict[str, set[str]] = {}

    dish_cal: dict[str, list[float]] = {}
    dish_exit: dict[str, list[float]] = {}
    dish_prot: dict[str, list[float]] = {}
    dish_fat: dict[str, list[float]] = {}
    dish_carb: dict[str, list[float]] = {}

    files_scanned = 0
    schools: set[str] = set()

    for path in _iter_reference_files(base_dir):
        if max_files is not None and files_scanned >= max_files:
            break
        try:
            rows, _ = parse_menu_file(path)
        except Exception:
            continue
        schools.add(path.parent.parent.name)
        files_scanned += 1
        for row in rows:
            sec = _norm_section(row.section)
            if not sec:
                continue
            section_exit.setdefault(sec, [])
            section_cal.setdefault(sec, [])
            section_prot.setdefault(sec, [])
            section_fat.setdefault(sec, [])
            section_carb.setdefault(sec, [])
            section_meals.setdefault(sec, set())

            if row.exit_g is not None and row.exit_g > 0:
                section_exit[sec].append(row.exit_g)
            if row.calories is not None and row.calories > 0:
                section_cal[sec].append(row.calories)
            if row.protein is not None:
                section_prot[sec].append(row.protein)
            if row.fat is not None:
                section_fat[sec].append(row.fat)
            if row.carb is not None:
                section_carb[sec].append(row.carb)
            if row.meal:
                section_meals[sec].add(row.meal)

            dish_key = _norm_dish(row.name)
            if dish_key and dish_key != "—":
                if row.calories is not None and row.calories > 0:
                    dish_cal.setdefault(dish_key, []).append(row.calories)
                if row.exit_g is not None and row.exit_g > 0:
                    dish_exit.setdefault(dish_key, []).append(row.exit_g)
                if row.protein is not None:
                    dish_prot.setdefault(dish_key, []).append(row.protein)
                if row.fat is not None:
                    dish_fat.setdefault(dish_key, []).append(row.fat)
                if row.carb is not None:
                    dish_carb.setdefault(dish_key, []).append(row.carb)

    sections: dict[str, PracticalSectionRule] = {}
    for sec, exits in section_exit.items():
        if len(exits) < MIN_PRACTICAL_SAMPLES:
            continue
        # Реальные мин/макс из эталонных файлов /food (без усечения)
        exit_min = round(min(exits), 1)
        exit_max = round(max(exits), 1)
        sections[sec] = PracticalSectionRule(
            section=sec,
            exit_min=exit_min,
            exit_max=exit_max,
            exit_avg=round(statistics.mean(exits), 1),
            cal_avg=_avg(section_cal.get(sec, [])),
            cal_min=round(min(section_cal[sec]), 1) if section_cal.get(sec) else None,
            prot_avg=_avg(section_prot.get(sec, [])),
            fat_avg=_avg(section_fat.get(sec, [])),
            carb_avg=_avg(section_carb.get(sec, [])),
            meals=sorted(section_meals.get(sec, set())),
            samples=len(exits),
        )

    dishes: dict[str, DishNorm] = {}
    for name, cals in dish_cal.items():
        if len(cals) < 3:
            continue
        exits = dish_exit.get(name, [])
        dishes[name] = DishNorm(
            name=name,
            cal_avg=round(statistics.mean(cals), 1),
            cal_min=round(min(cals), 1),
            exit_avg=round(statistics.mean(exits), 1) if exits else 0.0,
            prot_avg=_avg(dish_prot.get(name, [])),
            fat_avg=_avg(dish_fat.get(name, [])),
            carb_avg=_avg(dish_carb.get(name, [])),
            samples=len(cals),
        )

    return PracticalRules(
        sections=sections,
        dishes=dishes,
        built_at=datetime.now().isoformat(timespec="seconds"),
        files_scanned=files_scanned,
        schools_scanned=len(schools),
    )


def save_practical_rules(rules: PracticalRules, path: Path = CACHE_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "built_at": rules.built_at,
        "files_scanned": rules.files_scanned,
        "schools_scanned": rules.schools_scanned,
        "sections": {k: asdict(v) for k, v in rules.sections.items()},
        "dishes": {k: asdict(v) for k, v in rules.dishes.items()},
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_practical_rules(path: Path = CACHE_PATH) -> Optional[PracticalRules]:
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None

    sections: dict[str, PracticalSectionRule] = {}
    for k, v in data.get("sections", {}).items():
        payload = dict(v)
        if "cal_min" not in payload:
            payload["cal_min"] = payload.get("cal_avg")
        sections[k] = PracticalSectionRule(**payload)

    dishes: dict[str, DishNorm] = {}
    for k, v in data.get("dishes", {}).items():
        payload = dict(v)
        if "cal_min" not in payload:
            payload["cal_min"] = payload.get("cal_avg", 0.0)
        dishes[k] = DishNorm(**payload)

    return PracticalRules(
        sections=sections,
        dishes=dishes,
        built_at=data.get("built_at", ""),
        files_scanned=int(data.get("files_scanned", 0)),
        schools_scanned=int(data.get("schools_scanned", 0)),
    )


_rules_cache: Optional[PracticalRules] = None
_rules_lock_ts: float = 0.0


def get_practical_rules(force_rebuild: bool = False) -> PracticalRules:
    """Возвращает кэшированные практические правила; при отсутствии — строит."""
    global _rules_cache, _rules_lock_ts
    if not force_rebuild and _rules_cache is not None:
        return _rules_cache
    if not force_rebuild:
        loaded = load_practical_rules()
        if loaded is not None and loaded.sections:
            _rules_cache = loaded
            return loaded
    # защита от параллельной тяжёлой перестройки
    now = time.time()
    if now - _rules_lock_ts < 2 and _rules_cache is not None:
        return _rules_cache
    _rules_lock_ts = now
    rules = build_practical_rules()
    save_practical_rules(rules)
    _rules_cache = rules
    return rules


# ---------------------------------------------------------------------------
# Проверка
# ---------------------------------------------------------------------------

def _kbju_rescale(
    protein: Optional[float],
    fat: Optional[float],
    carb: Optional[float],
    calories: Optional[float],
    old_exit: Optional[float],
    new_exit: float,
) -> Optional[dict[str, float]]:
    if old_exit is None or old_exit <= 0:
        return None
    ratio = new_exit / old_exit
    out: dict[str, float] = {"Выход (г)": round(new_exit, 1)}
    if protein is not None:
        out["Белки"] = round(protein * ratio, 2)
    if fat is not None:
        out["Жиры"] = round(fat * ratio, 2)
    if carb is not None:
        out["Углеводы"] = round(carb * ratio, 2)
    if calories is not None:
        out["Калорийность"] = round(calories * ratio, 1)
    return out


def _match_section_eq(section: str, expected: str) -> bool:
    return _norm_section(section) == _norm_section(expected)


def _match_section_startswith(section: str, prefix: str) -> bool:
    return _norm_section(section).startswith(_norm_section(prefix))


def _is_drink_or_sweet_lunch(row: MenuRow) -> bool:
    sec = _norm_section(row.section)
    meal = _norm_meal(row.meal)
    if sec == "сладкое" and meal == "обед":
        return True
    if sec in {"напиток", "гор.напиток"}:
        return True
    return False


def _lookup_calorie_norm(row: MenuRow, rules: PracticalRules) -> Optional[float]:
    """Норма калорийности = минимальное значение из эталонных файлов для блюда/раздела."""
    dish = rules.dishes.get(_norm_dish(row.name))
    if dish is not None:
        return dish.cal_min
    sec = rules.sections.get(_norm_section(row.section))
    if sec is not None and sec.cal_min is not None:
        return sec.cal_min
    return None


def validate_row(row: MenuRow, rules: PracticalRules) -> list[Violation]:
    violations: list[Violation] = []
    sec_key = _norm_section(row.section)
    practical = rules.sections.get(sec_key)

    # --- Практические правила (приоритет) ---
    if practical is not None and row.exit_g is not None and row.exit_g > 0:
        if row.exit_g < practical.exit_min:
            rec = practical.exit_min
            violations.append(
                Violation(
                    rule_id="P-EXIT-MIN",
                    rule_source="practical",
                    rule_title=(
                        f"Раздел «{row.section}»: выход ниже эталонного минимума "
                        f"({practical.exit_min:g} г, n={practical.samples})"
                    ),
                    row_number=row.row_number,
                    dish_name=row.name,
                    section=row.section,
                    meal=row.meal,
                    field="Выход (г)",
                    current_value=row.exit_g,
                    recommended_value=rec,
                    suggestion=f"Увеличить выход с {row.exit_g:g} г до {rec:g} г",
                    kbju_suggestion=_kbju_rescale(
                        row.protein, row.fat, row.carb, row.calories, row.exit_g, rec
                    ),
                )
            )
        elif row.exit_g > practical.exit_max:
            rec = practical.exit_max
            violations.append(
                Violation(
                    rule_id="P-EXIT-MAX",
                    rule_source="practical",
                    rule_title=(
                        f"Раздел «{row.section}»: выход выше эталонного максимума "
                        f"({practical.exit_max:g} г, n={practical.samples})"
                    ),
                    row_number=row.row_number,
                    dish_name=row.name,
                    section=row.section,
                    meal=row.meal,
                    field="Выход (г)",
                    current_value=row.exit_g,
                    recommended_value=rec,
                    suggestion=f"Уменьшить выход с {row.exit_g:g} г до {rec:g} г",
                    kbju_suggestion=_kbju_rescale(
                        row.protein, row.fat, row.carb, row.calories, row.exit_g, rec
                    ),
                )
            )

    # Калорийность vs норма блюда/раздела — S6 и S7 отдельно.
    # Учитываем разброс рецептур: мелкие отклонения (276 vs 287, 118.75 vs 118.8) не ошибка.
    cal_norm = _lookup_calorie_norm(row, rules)
    if cal_norm is not None and row.calories is not None and _is_calorie_too_low(row.calories, cal_norm):
        threshold = round(cal_norm * CALORIE_REL_FLOOR, 1)
        for rule_id, title in (
            ("S6", "Калорийность существенно ниже нормы (сценарий 6)"),
            ("S7", "Калорийность существенно ниже нормы (сценарий 7)"),
        ):
            violations.append(
                Violation(
                    rule_id=rule_id,
                    rule_source="sanpin" if practical is None else "practical",
                    rule_title=f"{title}: эталон от {cal_norm:g} ккал",
                    row_number=row.row_number,
                    dish_name=row.name,
                    section=row.section,
                    meal=row.meal,
                    field="Калорийность",
                    current_value=row.calories,
                    recommended_value=cal_norm,
                    suggestion=(
                        f"Калорийность {row.calories:g} ккал заметно ниже эталона "
                        f"({cal_norm:g}). Рекомендуется не ниже ≈ {threshold:g} ккал"
                    ),
                )
            )

    # S9 — сумма БЖУ не должна превышать вес (отдельное универсальное правило)
    if (
        row.exit_g is not None
        and row.protein is not None
        and row.fat is not None
        and row.carb is not None
    ):
        bju_sum = row.protein + row.fat + row.carb
        if bju_sum > row.exit_g:
            # пропорционально ужать БЖУ к выходу
            ratio = row.exit_g / bju_sum if bju_sum else 1.0
            kbju = {
                "Белки": round(row.protein * ratio, 2),
                "Жиры": round(row.fat * ratio, 2),
                "Углеводы": round(row.carb * ratio, 2),
            }
            violations.append(
                Violation(
                    rule_id="S9",
                    rule_source="sanpin",
                    rule_title="Сумма Белки+Жиры+Углеводы больше выхода блюда",
                    row_number=row.row_number,
                    dish_name=row.name,
                    section=row.section,
                    meal=row.meal,
                    field="Белки+Жиры+Углеводы",
                    current_value=round(bju_sum, 2),
                    recommended_value=row.exit_g,
                    suggestion=(
                        f"Сумма БЖУ ({bju_sum:.2f} г) превышает выход ({row.exit_g:g} г). "
                        f"Пересчитать БЖУ пропорционально выходу."
                    ),
                    kbju_suggestion=kbju,
                )
            )

    # --- СанПиН fallback: только если нет практического правила по разделу ---
    if practical is not None:
        return violations

    exit_g = row.exit_g
    if exit_g is None:
        return violations

    for spec in SANPIN_FALLBACK_RULES:
        matched = False

        if spec.get("special") == "drink_or_sweet_lunch":
            matched = _is_drink_or_sweet_lunch(row) and exit_g < spec["exit_lt"]
        elif "section_eq" in spec:
            matched = (
                _match_section_eq(row.section, spec["section_eq"])
                and exit_g > spec.get("exit_gt", -1)
                and exit_g < spec["exit_lt"]
            )
        elif "section_startswith" in spec:
            matched = (
                _match_section_startswith(row.section, spec["section_startswith"])
                and exit_g > spec.get("exit_gt", -1)
                and exit_g < spec["exit_lt"]
            )
        elif spec.get("global_min"):
            # S10–S15 — отдельные сценарии, не объединяем
            matched = exit_g > 0 and exit_g < spec["exit_lt"]
        else:
            continue

        if not matched:
            continue

        rec = float(spec["recommend"])
        violations.append(
            Violation(
                rule_id=spec["id"],
                rule_source="sanpin",
                rule_title=spec["title"],
                row_number=row.row_number,
                dish_name=row.name,
                section=row.section or "—",
                meal=row.meal,
                field=spec["field"],
                current_value=exit_g,
                recommended_value=rec,
                suggestion=f"Увеличить выход с {exit_g:g} г до {rec:g} г",
                kbju_suggestion=_kbju_rescale(
                    row.protein, row.fat, row.carb, row.calories, exit_g, rec
                ),
            )
        )

    return violations


def validate_file(path: Path, rules: Optional[PracticalRules] = None) -> FileCheckResult:
    rules = rules or get_practical_rules()
    filename = path.name
    try:
        rows, file_type = parse_menu_file(path)
    except Exception as exc:
        return FileCheckResult(
            filename=filename,
            file_type=detect_file_type(filename),
            rows_checked=0,
            violations=[],
            error=str(exc),
        )

    violations: list[Violation] = []
    for row in rows:
        violations.extend(validate_row(row, rules))

    return FileCheckResult(
        filename=filename,
        file_type=file_type,
        rows_checked=len(rows),
        violations=violations,
        error=None,
    )


def validate_files(paths: list[Path], rules: Optional[PracticalRules] = None) -> dict[str, Any]:
    rules = rules or get_practical_rules()
    results = [validate_file(p, rules) for p in paths]
    total_violations = sum(len(r.violations) for r in results)
    return {
        "ok": True,
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "practical_rules": {
            "built_at": rules.built_at,
            "files_scanned": rules.files_scanned,
            "schools_scanned": rules.schools_scanned,
            "sections_count": len(rules.sections),
            "dishes_count": len(rules.dishes),
            "sections": {
                k: {
                    "exit_min": v.exit_min,
                    "exit_max": v.exit_max,
                    "exit_avg": v.exit_avg,
                    "cal_avg": v.cal_avg,
                    "samples": v.samples,
                }
                for k, v in sorted(rules.sections.items())
            },
        },
        "summary": {
            "files": len(results),
            "rows_checked": sum(r.rows_checked for r in results),
            "violations": total_violations,
            "files_with_errors": sum(1 for r in results if r.violations or r.error),
            "files_ok": sum(1 for r in results if not r.violations and not r.error),
        },
        "files": [_file_result_to_dict(r) for r in results],
    }


def _file_result_to_dict(result: FileCheckResult) -> dict[str, Any]:
    return {
        "filename": result.filename,
        "file_type": result.file_type,
        "rows_checked": result.rows_checked,
        "error": result.error,
        "violations_count": len(result.violations),
        "violations": [asdict(v) for v in result.violations],
    }


def make_temp_session_dir() -> Path:
    TEMP_CHECK_DIR.mkdir(parents=True, exist_ok=True)
    session = TEMP_CHECK_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
    session.mkdir(parents=True, exist_ok=True)
    return session


def cleanup_old_temp(max_age_hours: int = 6) -> None:
    if not TEMP_CHECK_DIR.exists():
        return
    cutoff = time.time() - max_age_hours * 3600
    for item in TEMP_CHECK_DIR.iterdir():
        try:
            if item.is_dir() and item.stat().st_mtime < cutoff:
                for child in item.iterdir():
                    child.unlink(missing_ok=True)
                item.rmdir()
        except OSError:
            pass
