"""ИИ-помощник справочного центра через OpenRouter."""

from __future__ import annotations

import json
import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from cachetools import TTLCache
from dotenv import load_dotenv

# Всегда грузим .env рядом с проектом (не зависеть от cwd systemd)
_BASE_DIR = Path(__file__).resolve().parent
load_dotenv(_BASE_DIR / ".env", override=False)

logger = logging.getLogger(__name__)

OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
MAX_DOC_CHARS = 70_000
MAX_PREVIEW_CHARS = 1_400
MAX_CATALOG_CHARS = 90_000

_text_cache: TTLCache = TTLCache(maxsize=128, ttl=3600)


def _api_key() -> str:
    return (os.getenv("OPENROUTER_API_KEY") or "").strip()


def _model() -> str:
    return (os.getenv("OPENROUTER_MODEL") or "openai/gpt-4o-mini").strip()


def _site_url() -> str:
    return (os.getenv("OPENROUTER_SITE_URL") or "https://openrouter.ai").strip()


def _site_name() -> str:
    return (os.getenv("OPENROUTER_SITE_NAME") or "FoodMonitoring").strip()


def _http_proxy() -> Optional[str]:
    return (
        os.getenv("OPENROUTER_HTTP_PROXY")
        or os.getenv("HTTPS_PROXY")
        or os.getenv("HTTP_PROXY")
        or ""
    ).strip() or None


def is_configured() -> bool:
    return bool(_api_key())


def extract_text_from_file(file_path: Path) -> str:
    """Извлекает текст из PDF / DOCX / TXT / таблиц."""
    if not file_path.exists() or not file_path.is_file():
        return ""

    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            parts: List[str] = []
            for page in reader.pages:
                try:
                    parts.append(page.extract_text() or "")
                except Exception:
                    continue
            return "\n".join(parts).strip()

        if ext == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs if p.text).strip()

        if ext in {".txt", ".md", ".csv", ".log"}:
            for encoding in ("utf-8", "cp1251", "latin-1"):
                try:
                    return file_path.read_text(encoding=encoding).strip()
                except UnicodeDecodeError:
                    continue
            return ""

        if ext in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            wb = load_workbook(str(file_path), read_only=True, data_only=True)
            rows: List[str] = []
            for sheet in wb.worksheets[:3]:
                rows.append(f"[{sheet.title}]")
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i > 200:
                        break
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append(" | ".join(cells))
            return "\n".join(rows).strip()

    except Exception as exc:
        logger.warning("Не удалось извлечь текст из %s: %s", file_path, exc)
        return ""

    return ""


def get_document_text(base_dir: Path, relative_path: str, doc_id: int) -> str:
    path = base_dir / relative_path
    try:
        mtime = path.stat().st_mtime if path.exists() else 0
    except OSError:
        mtime = 0
    cache_key = f"{doc_id}:{mtime}"
    if cache_key in _text_cache:
        return _text_cache[cache_key]

    text = extract_text_from_file(path)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    _text_cache[cache_key] = text
    return text


def _format_openrouter_error(status: int, body: str) -> str:
    hint = ""
    lower = (body or "").lower()
    if status == 403:
        if "security policy" in lower or "cloudflare" in lower or "access denied" in lower:
            hint = (
                " Сервер, скорее всего, заблокирован Cloudflare/по региону. "
                "Проверьте с сервера curl к OpenRouter или задайте OPENROUTER_HTTP_PROXY."
            )
        else:
            hint = (
                " Проверьте ключ OPENROUTER_API_KEY и баланс на openrouter.ai; "
                "для продакшена укажите реальный OPENROUTER_SITE_URL (не localhost)."
            )
    elif status == 401:
        hint = " Неверный или отозванный OPENROUTER_API_KEY."
    elif status == 402:
        hint = " Недостаточно средств на аккаунте OpenRouter."

    snippet = re.sub(r"\s+", " ", (body or "").strip())[:220]
    if snippet:
        return f"OpenRouter вернул ошибку {status}: {snippet}.{hint}"
    return f"OpenRouter вернул ошибку {status}.{hint}"


async def openrouter_chat(
    messages: List[Dict[str, str]],
    *,
    temperature: float = 0.3,
    max_tokens: int = 1800,
    response_format: Optional[Dict[str, str]] = None,
) -> str:
    key = _api_key()
    if not key:
        raise RuntimeError("OPENROUTER_API_KEY не задан")

    site_url = _site_url()
    site_name = _site_name()
    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "HTTP-Referer": site_url,
        "Referer": site_url,
        "X-Title": site_name,
        "X-OpenRouter-Title": site_name,
        "User-Agent": f"FoodMonitoring/1.0 ({site_name})",
    }
    payload: Dict[str, Any] = {
        "model": _model(),
        "messages": messages,
        "temperature": temperature,
        "max_tokens": max_tokens,
    }
    if response_format:
        payload["response_format"] = response_format

    proxy = _http_proxy()
    client_kwargs: Dict[str, Any] = {"timeout": 90.0}
    if proxy:
        client_kwargs["proxy"] = proxy

    async with httpx.AsyncClient(**client_kwargs) as client:
        resp = await client.post(OPENROUTER_URL, headers=headers, json=payload)
        if resp.status_code >= 400:
            detail = resp.text[:800]
            logger.error("OpenRouter error %s: %s", resp.status_code, detail)
            raise RuntimeError(_format_openrouter_error(resp.status_code, detail))
        data = resp.json()

    try:
        return data["choices"][0]["message"]["content"].strip()
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("Некорректный ответ OpenRouter") from exc


def _parse_json_payload(raw: str) -> Dict[str, Any]:
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{[\s\S]*\}", raw)
        if match:
            try:
                return json.loads(match.group(0))
            except json.JSONDecodeError:
                pass
    return {"reply": raw, "documents": []}


def build_catalog_entries(
    documents: List[Any],
    base_dir: Path,
    document_types: Dict[str, str],
    categories: Dict[int, str],
    query: str = "",
) -> List[Dict[str, Any]]:
    tokens = [t for t in re.split(r"\W+", query.lower()) if len(t) > 2]
    entries: List[Dict[str, Any]] = []
    total_chars = 0

    scored = []
    for doc in documents:
        hay = " ".join(
            filter(
                None,
                [
                    doc.title or "",
                    doc.description or "",
                    doc.tags or "",
                    document_types.get(doc.document_type, ""),
                    categories.get(doc.category_id or 0, ""),
                ],
            )
        ).lower()
        score = sum(1 for t in tokens if t in hay) if tokens else 0
        scored.append((score, doc))

    scored.sort(key=lambda x: (-x[0], -(x[1].downloads_count or 0), -(x[1].id or 0)))

    for score, doc in scored:
        include_preview = score > 0 or not tokens or len(entries) < 12
        preview = ""
        if include_preview and (doc.file_extension or "").lower() in {
            ".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".xlsm"
        }:
            text = get_document_text(base_dir, doc.file_path, doc.id)
            preview = text[:MAX_PREVIEW_CHARS]

        entry = {
            "id": doc.id,
            "title": doc.title,
            "description": (doc.description or "")[:500],
            "type": document_types.get(doc.document_type, doc.document_type),
            "category": categories.get(doc.category_id or 0, "Без категории"),
            "tags": doc.tags or "",
            "extension": doc.file_extension or "",
            "downloads": doc.downloads_count or 0,
            "preview": preview,
        }
        chunk = json.dumps(entry, ensure_ascii=False)
        if total_chars + len(chunk) > MAX_CATALOG_CHARS and entries:
            break
        entries.append(entry)
        total_chars += len(chunk)

    return entries


async def search_documents_with_ai(
    *,
    user_message: str,
    documents: List[Any],
    base_dir: Path,
    document_types: Dict[str, str],
    categories: Dict[int, str],
    history: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    catalog = build_catalog_entries(
        documents, base_dir, document_types, categories, user_message
    )
    catalog_json = json.dumps(catalog, ensure_ascii=False)

    system = (
        "Ты — ИИ-помощник справочного центра по организации питания в образовательных учреждениях. "
        "По запросу пользователя подбери подходящие документы из каталога. "
        "Отвечай только на русском. Не выдумывай документы вне каталога. "
        "Если подходящих нет — честно скажи и предложи уточнить запрос. "
        "Верни строго JSON объект вида:\n"
        '{"reply":"понятный ответ пользователю (можно с markdown)","documents":[{"id":1,"reason":"почему подходит"}]}\n'
        "В documents включай только реально существующие id из каталога, максимум 5."
    )

    messages: List[Dict[str, str]] = [{"role": "system", "content": system}]
    if history:
        for item in history[-6:]:
            role = item.get("role")
            content = (item.get("content") or "").strip()
            if role in {"user", "assistant"} and content:
                messages.append({"role": role, "content": content[:2000]})

    messages.append(
        {
            "role": "user",
            "content": (
                f"Каталог документов:\n{catalog_json}\n\n"
                f"Запрос пользователя: {user_message.strip()}"
            ),
        }
    )

    raw = await openrouter_chat(
        messages,
        temperature=0.2,
        max_tokens=1600,
        response_format={"type": "json_object"},
    )
    parsed = _parse_json_payload(raw)

    id_map = {d.id: d for d in documents}
    recommended = []
    for item in parsed.get("documents") or []:
        try:
            doc_id = int(item.get("id"))
        except (TypeError, ValueError, AttributeError):
            continue
        doc = id_map.get(doc_id)
        if not doc:
            continue
        recommended.append(
            {
                "id": doc.id,
                "title": doc.title,
                "type": document_types.get(doc.document_type, "Документ"),
                "url": f"/knowledge-base/document/{doc.id}",
                "cover_url": f"/knowledge-base/cover/{doc.id}" if doc.cover_image_path else None,
                "reason": (item.get("reason") or "").strip()[:300],
            }
        )

    reply = (parsed.get("reply") or "").strip() or "Не удалось сформировать ответ. Попробуйте переформулировать запрос."
    return {"reply": reply, "documents": recommended}


async def ask_about_document(
    *,
    document: Any,
    base_dir: Path,
    document_types: Dict[str, str],
    question: Optional[str] = None,
    history: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    text = get_document_text(base_dir, document.file_path, document.id)
    meta = (
        f"Название: {document.title}\n"
        f"Тип: {document_types.get(document.document_type, document.document_type)}\n"
        f"Описание: {document.description or '—'}\n"
        f"Теги: {document.tags or '—'}\n"
        f"Формат: {document.file_extension or '—'}"
    )

    if not text:
        text_block = (
            "Текст файла извлечь не удалось (возможно, это скан, видео или защищённый файл). "
            "Опирайся только на метаданные выше и честно скажи, если данных недостаточно."
        )
    else:
        text_block = text[:MAX_DOC_CHARS]

    q = (question or "").strip()
    if not q:
        q = (
            "Сделай краткую выжимку документа: основные положения, важные требования, "
            "сроки/нормы если есть, и на что обратить внимание практическому специалисту."
        )

    system = (
        "Ты — ИИ-помощник справочного центра. Отвечай на русском, ясно и по делу. "
        "Опирайся только на предоставленный текст документа и метаданные. "
        "Если ответа в тексте нет — так и скажи. Не выдумывай факты. "
        "Форматируй ответ структурированно (короткие абзацы или маркированные списки)."
    )

    messages: List[Dict[str, str]] = [{"role": "system", "content": system}]
    if history:
        for item in history[-8:]:
            role = item.get("role")
            content = (item.get("content") or "").strip()
            if role in {"user", "assistant"} and content:
                messages.append({"role": role, "content": content[:3000]})

    messages.append(
        {
            "role": "user",
            "content": (
                f"{meta}\n\n--- Текст документа ---\n{text_block}\n--- Конец ---\n\n"
                f"Вопрос пользователя: {q}"
            ),
        }
    )

    reply = await openrouter_chat(messages, temperature=0.25, max_tokens=2200)
    return {
        "reply": reply,
        "has_text": bool(text),
        "chars_used": min(len(text), MAX_DOC_CHARS) if text else 0,
    }
